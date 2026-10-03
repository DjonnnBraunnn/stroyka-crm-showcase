"""Отчёты директору. Восемь отчётов из ТЗ, каждый за день / неделю / месяц,
плюс сводная панель на главной. Любой отчёт выгружается в Excel."""
from datetime import date, timedelta, datetime
from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session
from ..database import get_db
from ..auth import Me, get_current_user
from .. import models as m
from .. import serializers as s
from .camp import structure as camp_structure

router = APIRouter(prefix="/api/reports", tags=["reports"])


def bounds(period: str) -> tuple[date, date]:
    today = date.today()
    if period == "day":
        return today, today
    if period == "week":
        return today - timedelta(days=6), today
    if period == "year":
        return today - timedelta(days=365), today
    return today - timedelta(days=29), today


# ── 1. Состояние объектов ────────────────────────────────────────────────

def report_sites(db, cid: int, period: str):
    from .sites import site_stats
    sites = db.query(m.Site).filter(m.Site.company_id == cid).order_by(m.Site.direction_id, m.Site.name).all()
    rows = [s.site(x, site_stats(db, x)) for x in sites]
    by_status = {}
    for r in rows:
        by_status[r["status_name"]] = by_status.get(r["status_name"], 0) + 1
    return {"rows": rows, "by_status": by_status,
            "late": sum(1 for r in rows if r["is_late"]),
            "over_budget": sum(1 for r in rows if r["over_budget"])}


# ── 2. Выполнение плана ──────────────────────────────────────────────────

def report_plan(db, cid: int, period: str):
    d1, d2 = bounds(period)
    tasks = db.query(m.Task).filter(m.Task.company_id == cid).all()
    in_period = [t for t in tasks if t.deadline and d1 <= t.deadline <= d2]
    done_in = [t for t in tasks if t.completed_at and d1 <= t.completed_at.date() <= d2]
    today = date.today()
    sites = db.query(m.Site).filter(m.Site.company_id == cid, m.Site.status == "active").all()
    rows = []
    for x in sites:
        st = [t for t in tasks if t.site_id == x.id]
        planned = [t for t in st if t.deadline and d1 <= t.deadline <= d2]
        done = [t for t in planned if t.status == "done"]
        # ожидаемый % по календарю: сколько времени прошло от начала до плановой сдачи
        expected = None
        if x.start_plan and x.end_plan and x.end_plan > x.start_plan:
            expected = max(0, min(100, round((today - x.start_plan).days / (x.end_plan - x.start_plan).days * 100)))
        rows.append({"site_id": x.id, "site_name": x.name, "direction": s.direction(x.direction),
                     "progress": x.progress, "expected": expected,
                     "lag": (expected - x.progress) if expected is not None else None,
                     "planned": len(planned), "done": len(done),
                     "pct": round(len(done) / len(planned) * 100) if planned else None})
    return {"from": d1.isoformat(), "to": d2.isoformat(),
            "tasks_planned": len(in_period), "tasks_done": len(done_in),
            "pct": round(sum(1 for t in in_period if t.status == "done") / len(in_period) * 100) if in_period else 0,
            "rows": rows}


# ── 3. Выработка техники ─────────────────────────────────────────────────

def report_equipment(db, cid: int, period: str):
    d1, d2 = bounds(period)
    eq = db.query(m.Equipment).filter(m.Equipment.company_id == cid).all()
    logs = db.query(m.WorkLog).filter(m.WorkLog.company_id == cid, m.WorkLog.date >= d1, m.WorkLog.date <= d2).all()
    days = (d2 - d1).days + 1
    rows = []
    for e in eq:
        el = [w for w in logs if w.equipment_id == e.id]
        hours = sum(w.hours or 0 for w in el)
        rows.append({"id": e.id, "name": e.name, "kind": e.kind, "plate": e.plate, "status": e.status,
                     "status_name": m.EQUIPMENT_STATUSES.get(e.status), "site_name": e.site.name if e.site else None,
                     "shifts": len(el), "hours": round(hours, 1),
                     "output": round(sum(w.output or 0 for w in el), 1),
                     "output_unit": el[0].output_unit if el else "",
                     "utilization": round(hours / (days * 10) * 100) if days else 0,   # смена 10 ч
                     "fuel": round(sum(w.fuel_used or 0 for w in el))})
    rows.sort(key=lambda r: -r["hours"])
    by_status = {}
    for e in eq:
        n = m.EQUIPMENT_STATUSES.get(e.status, e.status)
        by_status[n] = by_status.get(n, 0) + 1
    return {"from": d1.isoformat(), "to": d2.isoformat(), "rows": rows, "by_status": by_status,
            "total_hours": round(sum(r["hours"] for r in rows), 1),
            "avg_utilization": round(sum(r["utilization"] for r in rows) / len(rows)) if rows else 0,
            "idle": sum(1 for e in eq if e.status in ("idle", "repair"))}


# ── 4. Расход ГСМ ────────────────────────────────────────────────────────

def report_fuel(db, cid: int, period: str):
    d1, d2 = bounds(period)
    logs = db.query(m.WorkLog).filter(m.WorkLog.company_id == cid, m.WorkLog.date >= d1, m.WorkLog.date <= d2).all()
    by_eq: dict[int, dict] = {}
    by_day: dict[str, float] = {}
    for w in logs:
        e = w.equipment
        r = by_eq.setdefault(e.id, {"id": e.id, "name": e.name, "plate": e.plate, "hours": 0,
                                    "issued": 0, "used": 0, "norm": 0})
        r["hours"] += w.hours or 0
        r["issued"] += w.fuel_issued or 0
        r["used"] += w.fuel_used or 0
        r["norm"] += (w.hours or 0) * (e.fuel_norm or 0)
        k = w.date.isoformat()
        by_day[k] = by_day.get(k, 0) + (w.fuel_used or 0)
    rows = []
    for r in by_eq.values():
        over = r["used"] - r["norm"]
        rows.append({**r, "hours": round(r["hours"], 1), "issued": round(r["issued"]), "used": round(r["used"]),
                     "norm": round(r["norm"]), "over": round(over),
                     "over_pct": round(over / r["norm"] * 100) if r["norm"] else 0,
                     "per_hour": round(r["used"] / r["hours"], 1) if r["hours"] else 0})
    rows.sort(key=lambda r: -r["over"])
    fuel_cost = db.query(func.coalesce(func.sum(m.Expense.amount), 0)).filter(
        m.Expense.company_id == cid, m.Expense.category == "fuel",
        m.Expense.date >= d1, m.Expense.date <= d2).scalar()
    return {"from": d1.isoformat(), "to": d2.isoformat(), "rows": rows,
            "total_used": round(sum(r["used"] for r in rows)), "total_norm": round(sum(r["norm"] for r in rows)),
            "total_over": round(sum(r["over"] for r in rows)),
            "fuel_cost": fuel_cost,
            "by_day": [{"date": k, "used": round(v)} for k, v in sorted(by_day.items())]}


# ── 5. Численность персонала ─────────────────────────────────────────────

def report_staff(db, cid: int, period: str):
    today = date.today()
    emps = db.query(m.Employee).filter(m.Employee.company_id == cid, m.Employee.status != "dismissed").all()
    rots = db.query(m.Rotation).filter(m.Rotation.company_id == cid, m.Rotation.date_from <= today,
                                       m.Rotation.date_to >= today).all()
    on_site = {r.employee_id: r for r in rots}
    by_site: dict[str, int] = {}
    for r in rots:
        by_site[r.site.name] = by_site.get(r.site.name, 0) + 1
    by_cat = {}
    for e in emps:
        n = m.EMPLOYEE_CATEGORIES.get(e.category, e.category)
        by_cat[n] = by_cat.get(n, 0) + 1
    by_dir = {}
    for e in emps:
        n = e.direction.short or e.direction.name if e.direction else "Без направления"
        by_dir[n] = by_dir.get(n, 0) + 1
    week = today + timedelta(days=7)
    ending = db.query(m.Rotation).filter(m.Rotation.company_id == cid, m.Rotation.date_to >= today,
                                         m.Rotation.date_to <= week).count()
    starting = db.query(m.Rotation).filter(m.Rotation.company_id == cid, m.Rotation.date_from > today,
                                           m.Rotation.date_from <= week).count()
    return {"total": len(emps), "on_site": len(on_site), "rest": len(emps) - len(on_site),
            "by_site": [{"site": k, "count": v} for k, v in sorted(by_site.items(), key=lambda x: -x[1])],
            "by_category": by_cat, "by_direction": by_dir,
            "rotation_ending_week": ending, "rotation_starting_week": starting,
            "sick": sum(1 for e in emps if e.status == "sick")}


# ── 6. Закрытие вакансий ─────────────────────────────────────────────────

def report_vacancies(db, cid: int, period: str):
    rows = db.query(m.Vacancy).filter(m.Vacancy.company_id == cid).order_by(m.Vacancy.status, m.Vacancy.opened_at).all()
    today = date.today()
    out = []
    for v in rows:
        out.append({**s.vacancy(v), "open_days": (today - v.opened_at).days if v.opened_at else 0,
                    "pct": round(v.closed / v.count * 100) if v.count else 0})
    need = sum(v.count - v.closed for v in rows if v.status == "open")
    return {"rows": out, "open": sum(1 for v in rows if v.status == "open"),
            "closed": sum(1 for v in rows if v.status == "closed"), "need_people": need,
            "avg_days_open": round(sum(r["open_days"] for r in out if r["status"] == "open") /
                                   max(1, sum(1 for r in out if r["status"] == "open")))}


# ── 7. Просроченные задачи ───────────────────────────────────────────────

def report_overdue(db, cid: int, period: str):
    today = date.today()
    tasks = db.query(m.Task).filter(m.Task.company_id == cid, m.Task.status != "done",
                                    m.Task.deadline < today).order_by(m.Task.deadline).all()
    rows = [{**s.task(t), "days_overdue": (today - t.deadline).days} for t in tasks]
    by_assignee: dict[str, int] = {}
    by_site: dict[str, int] = {}
    for r in rows:
        a = r["assignee"]["name"] if r["assignee"] else "Без ответственного"
        by_assignee[a] = by_assignee.get(a, 0) + 1
        by_site[r["site_name"]] = by_site.get(r["site_name"], 0) + 1
    soon = db.query(m.Task).filter(m.Task.company_id == cid, m.Task.status != "done",
                                   m.Task.deadline >= today, m.Task.deadline <= today + timedelta(days=3)).count()
    return {"rows": rows, "total": len(rows), "due_soon": soon,
            "by_assignee": [{"name": k, "count": v} for k, v in sorted(by_assignee.items(), key=lambda x: -x[1])],
            "by_site": [{"site": k, "count": v} for k, v in sorted(by_site.items(), key=lambda x: -x[1])],
            "critical": sum(1 for r in rows if r["priority"] == "critical")}


# ── 8. Финансовый отчёт ──────────────────────────────────────────────────

def report_finance(db, cid: int, period: str):
    from .finance import summary
    d1, d2 = bounds(period)
    exp = db.query(m.Expense).filter(m.Expense.company_id == cid, m.Expense.date >= d1, m.Expense.date <= d2).all()
    acts = db.query(m.Act).filter(m.Act.company_id == cid, m.Act.date >= d1, m.Act.date <= d2,
                                  m.Act.status.in_(["signed", "paid"])).all()
    pays = db.query(m.Payment).filter(m.Payment.company_id == cid, m.Payment.date >= d1, m.Payment.date <= d2).all()
    by_cat = {}
    for x in exp:
        n = m.EXPENSE_CATEGORIES.get(x.category, x.category)
        by_cat[n] = by_cat.get(n, 0) + (x.amount or 0)
    class _U:  # summary() ждёт Me; нам нужен только company_id
        company_id = cid
    sm = summary(_U(), db)
    return {"from": d1.isoformat(), "to": d2.isoformat(),
            "spent": sum(x.amount or 0 for x in exp), "acts": sum(a.amount or 0 for a in acts),
            "paid_in": sum(p.amount or 0 for p in pays if p.direction == "in"),
            "paid_out": sum(p.amount or 0 for p in pays if p.direction == "out"),
            "by_category": [{"name": k, "amount": v} for k, v in sorted(by_cat.items(), key=lambda x: -x[1])],
            "sites": sm["rows"], "totals": sm["totals"]}


REPORTS = {
    "sites": ("Состояние объектов", report_sites),
    "plan": ("Выполнение плана", report_plan),
    "equipment": ("Выработка техники", report_equipment),
    "fuel": ("Расход ГСМ", report_fuel),
    "staff": ("Численность персонала", report_staff),
    "vacancies": ("Закрытие вакансий", report_vacancies),
    "overdue": ("Просроченные задачи", report_overdue),
    "finance": ("Финансовый отчёт", report_finance),
}


@router.get("/dashboard")
def dashboard(user: Me = Depends(get_current_user), db: Session = Depends(get_db)):
    """Главная директора: всё важное за один запрос."""
    cid = user.company_id
    today = date.today()
    sites = report_sites(db, cid, "month")
    overdue = report_overdue(db, cid, "month")
    staff = report_staff(db, cid, "month")
    fuel = report_fuel(db, cid, "week")
    fuel30 = report_fuel(db, cid, "month")
    eq = report_equipment(db, cid, "week")
    fin = report_finance(db, cid, "month")
    vac = report_vacancies(db, cid, "month")
    camp = camp_structure(user, db)
    dirs = db.query(m.Direction).filter(m.Direction.company_id == cid).order_by(m.Direction.order).all()
    by_direction = []
    for d in dirs:
        ds = [r for r in sites["rows"] if r["direction_id"] == d.id]
        by_direction.append({
            "direction": s.direction(d), "sites": len(ds),
            "active": sum(1 for r in ds if r["status"] == "active"),
            "tasks_open": sum(r["tasks_open"] for r in ds),
            "tasks_overdue": sum(r["tasks_overdue"] for r in ds),
            "people": sum(r["people_today"] for r in ds),
            "spent": sum(r["spent"] for r in ds), "budget": sum(r["budget"] for r in ds),
            "progress": round(sum(r["progress"] for r in ds) / len(ds)) if ds else 0,
        })
    prod = db.query(func.coalesce(func.sum(m.Production.volume), 0)).filter(
        m.Production.company_id == cid, m.Production.date >= today - timedelta(days=6)).scalar()
    # В ленту директора — события компании; личные напоминания исполнителям
    # (ежедневные «просрочено») идут им в колокольчик и сюда не попадают.
    notes = db.query(m.Notification).filter(m.Notification.company_id == cid,
                                            m.Notification.user_id.is_(None)).order_by(
        m.Notification.created_at.desc()).limit(8).all()
    docs_review = db.query(m.Document).filter(m.Document.company_id == cid, m.Document.status == "review").count()
    return {
        "today": today.isoformat(),
        "sites": {"total": len(sites["rows"]), "active": sites["by_status"].get("В работе", 0),
                  "late": sites["late"], "over_budget": sites["over_budget"],
                  "rows": [r for r in sites["rows"] if r["status"] == "active"][:8]},
        "by_direction": by_direction,
        "tasks": {"overdue": overdue["total"], "due_soon": overdue["due_soon"], "critical": overdue["critical"],
                  "top": overdue["rows"][:6]},
        "staff": {"total": staff["total"], "on_site": staff["on_site"], "rest": staff["rest"],
                  "ending_week": staff["rotation_ending_week"], "starting_week": staff["rotation_starting_week"]},
        "camp": {"beds_total": camp["beds_total"], "occupied": camp["occupied"], "free": camp["free"],
                 "expected": camp["expected"], "load_pct": camp["load_pct"]},
        "equipment": {"total": len(eq["rows"]), "idle": eq["idle"], "hours_week": eq["total_hours"],
                      "utilization": eq["avg_utilization"], "by_status": eq["by_status"]},
        "fuel": {"used_week": fuel["total_used"], "norm_week": fuel["total_norm"], "over_week": fuel["total_over"],
                 "by_day": fuel["by_day"], "by_day_30": fuel30["by_day"],
                 "norm_per_day_30": round(fuel30["total_norm"] / max(1, len(fuel30["by_day"])))},
        "production_week": prod,
        "finance": {"spent_month": fin["spent"], "acts_month": fin["acts"], "paid_in_month": fin["paid_in"],
                    "receivable": fin["totals"]["receivable"], "over_budget": fin["totals"]["over_budget_count"],
                    "by_category": fin["by_category"][:6]},
        "vacancies": {"open": vac["open"], "need_people": vac["need_people"]},
        "docs_review": docs_review,
        "feed": [s.notification(n) for n in notes],
    }


@router.get("/{name}")
def get_report(name: str, period: str = "month", user: Me = Depends(get_current_user), db: Session = Depends(get_db)):
    if name not in REPORTS:
        raise HTTPException(404, "Нет такого отчёта")
    title, fn = REPORTS[name]
    return {"name": name, "title": title, "period": period, **fn(db, user.company_id, period)}


@router.get("/{name}/excel")
def report_excel(name: str, period: str = "month", user: Me = Depends(get_current_user), db: Session = Depends(get_db)):
    """Выгрузка табличной части отчёта в Excel: то, что бухгалтер несёт в 1С."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    if name not in REPORTS:
        raise HTTPException(404, "Нет такого отчёта")
    title, fn = REPORTS[name]
    data = fn(db, user.company_id, period)
    rows = data.get("rows") or data.get("sites") or []
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append([title, f"Период: {data.get('from', '')} — {data.get('to', date.today().isoformat())}"])
    ws["A1"].font = Font(bold=True, size=13)
    if rows:
        cols = [k for k, v in rows[0].items() if not isinstance(v, (dict, list))]
        ws.append(cols)
        for c in ws[2]:
            c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="FDE68A")
        for r in rows:
            ws.append([r.get(c) for c in cols])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 18
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"{name}_{period}_{date.today().isoformat()}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={fname}"})
